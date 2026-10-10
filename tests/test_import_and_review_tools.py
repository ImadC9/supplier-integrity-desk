"""Column mapping, CSV import, clusters, comparison and reports."""

from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import load_workbook

from supplier_agent.cli import main as cli_main
from supplier_agent.clusters import build_clusters, survivor_ranking
from supplier_agent.columns import apply_mapping, missing_required, suggest_mapping
from supplier_agent.compare import compare_records, mask
from supplier_agent.detector import DetectionConfig, DuplicateDetector, similar_pairs
from supplier_agent.io import load_supplier_master, read_supplier_table, template_workbook, validate_supplier_master
from supplier_agent.paths import SAMPLE_WORKBOOK
from supplier_agent.reports import build_excel_report
from supplier_agent.triage import scan_suppliers

EXAMPLE_CSV = Path(__file__).resolve().parents[1] / "examples" / "vendors_erp_export.csv"


class ColumnMappingTests(unittest.TestCase):
    def test_common_erp_headers_are_mapped(self):
        mapping = suggest_mapping(["LIFNR", "NAME1", "STCD1", "LAND1", "Street", "Zip", "IBAN", "Something else"])
        self.assertEqual(mapping["SUPPLIER_ID"], "LIFNR")
        self.assertEqual(mapping["SUPPLIER_NAME"], "NAME1")
        self.assertEqual(mapping["TAX_ID"], "STCD1")
        self.assertEqual(mapping["COUNTRY_CODE"], "LAND1")
        self.assertEqual(mapping["ADDRESS_1"], "Street")
        self.assertEqual(mapping["POSTAL_CODE"], "Zip")
        self.assertEqual(mapping["BANK_ACCOUNT"], "IBAN")
        self.assertIsNone(mapping["CITY"])

    def test_canonical_names_win_and_columns_are_used_once(self):
        mapping = suggest_mapping(["supplier_id", "Vendor ID", "Supplier Name", "Country", "Country Code"])
        self.assertEqual(mapping["SUPPLIER_ID"], "supplier_id")
        self.assertEqual(mapping["COUNTRY_CODE"], "Country Code")
        self.assertEqual(len([v for v in mapping.values() if v]), len({v for v in mapping.values() if v}))

    def test_required_fields_are_enforced_and_extras_kept(self):
        frame = pd.DataFrame({"Vendor": ["1"], "Notes": ["x"]})
        mapping = suggest_mapping(frame.columns)
        self.assertEqual(missing_required(mapping), ["SUPPLIER_NAME"])
        with self.assertRaisesRegex(ValueError, "Supplier name"):
            apply_mapping(frame, mapping)
        mapping["SUPPLIER_NAME"] = "Notes"
        mapped = apply_mapping(frame, mapping)
        self.assertEqual(mapped.loc[0, "SUPPLIER_NAME"], "x")
        self.assertEqual(mapped.loc[0, "TAX_ID"], "")

    def test_minimal_file_validates_with_blank_matching_columns(self):
        frame = validate_supplier_master(pd.DataFrame({"SUPPLIER_ID": ["a"], "SUPPLIER_NAME": ["Acme"]}))
        self.assertIn("BANK_ACCOUNT", frame.columns)
        self.assertEqual(frame.loc[0, "SUPPLIER_ID"], "A")


class FileImportTests(unittest.TestCase):
    def test_semicolon_csv_with_erp_headers_loads_and_scans(self):
        suppliers = load_supplier_master(EXAMPLE_CSV)
        self.assertGreater(len(suppliers), 100)
        result = scan_suppliers(suppliers)
        self.assertGreater(len(result.clusters), 0)
        self.assertIn("PLACEHOLDER_IDENTIFIER", set(result.quality_findings.RULE_CODE))
        self.assertIn("PHONE", result.absent_fields)
        self.assertNotIn("PHONE", set(result.quality_findings.FIELD))

    def test_csv_bytes_and_latin1_are_supported(self):
        raw = "Vendor No,Vendor Name,City\n1,Café Señor,Málaga\n".encode("cp1252")
        frame = read_supplier_table(raw)
        self.assertEqual(frame.loc[0, "Vendor Name"], "Café Señor")

    def test_label_sheet_cannot_be_selected(self):
        with self.assertRaisesRegex(ValueError, "Sheet not available"):
            read_supplier_table(SAMPLE_WORKBOOK, sheet="Ground_Truth")

    def test_template_contains_every_field(self):
        workbook = load_workbook(BytesIO(template_workbook()))
        self.assertEqual(workbook.sheetnames, ["Supplier_Master", "Field_Guide"])
        headers = [cell.value for cell in workbook["Supplier_Master"][1]]
        self.assertIn("SUPPLIER_ID", headers)
        self.assertIn("BANK_ACCOUNT", headers)


def _frame(rows):
    base = {
        "COUNTRY_CODE": "GB",
        "ADDRESS_1": "",
        "CITY": "",
        "POSTAL_CODE": "",
        "BANK_COUNTRY": "GB",
        "BANK_ACCOUNT": "",
        "TAX_ID": "",
        "ASSOCIATION_NO": "",
        "ADDRESS_2": "",
    }
    return validate_supplier_master(pd.DataFrame([base | row for row in rows]))


class MatchingTests(unittest.TestCase):
    def test_placeholder_identifiers_never_link_records(self):
        frame = _frame(
            [
                {"SUPPLIER_ID": "1", "SUPPLIER_NAME": "Alpha Marine", "TAX_ID": "N/A"},
                {"SUPPLIER_ID": "2", "SUPPLIER_NAME": "Zulu Bakery", "TAX_ID": "N/A"},
                {"SUPPLIER_ID": "3", "SUPPLIER_NAME": "Kilo Steel", "TAX_ID": "0000"},
                {"SUPPLIER_ID": "4", "SUPPLIER_NAME": "Hotel Paints", "TAX_ID": "0000"},
            ]
        )
        self.assertTrue(DuplicateDetector().detect(frame).empty)

    def test_widely_shared_identifier_is_ignored(self):
        rows = [{"SUPPLIER_ID": str(i), "SUPPLIER_NAME": f"Company {chr(65 + i)} {i * 7919}", "TAX_ID": "GB999"} for i in range(30)]
        pairs = DuplicateDetector(DetectionConfig(name_candidate_threshold=100, address_candidate_threshold=100)).detect(_frame(rows))
        self.assertFalse(pairs["SAME_TAX_ID"].any() if not pairs.empty else False)

    def test_records_without_country_match_only_when_allowed(self):
        frame = _frame(
            [
                {"SUPPLIER_ID": "1", "SUPPLIER_NAME": "Acme Trading Ltd", "COUNTRY_CODE": "", "TAX_ID": "X1"},
                {"SUPPLIER_ID": "2", "SUPPLIER_NAME": "ACME Trading Limited", "COUNTRY_CODE": "", "TAX_ID": "X1"},
            ]
        )
        self.assertTrue(DuplicateDetector().detect(frame).empty)
        pairs = DuplicateDetector(DetectionConfig(require_country=False)).detect(frame)
        self.assertEqual(len(pairs), 1)
        self.assertTrue(pairs.iloc[0]["SAME_TAX_ID"])

    def test_vectorized_similarity_matches_pairwise_definition(self):
        values = ["acme trading", "acme trading co", "", "zulu bakery", "acme"]
        self.assertEqual(similar_pairs(values, 72), {(0, 1), (0, 4), (1, 4)})


class ClusterTests(unittest.TestCase):
    def test_transitive_pairs_form_one_cluster_with_an_active_survivor(self):
        frame = _frame(
            [
                {"SUPPLIER_ID": "A", "SUPPLIER_NAME": "Acme Ltd", "TAX_ID": "T1", "ACTIVE_STATUS": "INACTIVE", "PO_COUNT_12M": "50"},
                {
                    "SUPPLIER_ID": "B",
                    "SUPPLIER_NAME": "Acme Limited",
                    "TAX_ID": "T1",
                    "BANK_ACCOUNT": "B1",
                    "ACTIVE_STATUS": "ACTIVE",
                    "PO_COUNT_12M": "3",
                },
                {"SUPPLIER_ID": "C", "SUPPLIER_NAME": "ACME", "BANK_ACCOUNT": "B1", "ACTIVE_STATUS": "ACTIVE", "PO_COUNT_12M": "9"},
                {"SUPPLIER_ID": "D", "SUPPLIER_NAME": "Unrelated Foods", "TAX_ID": "T9"},
            ]
        )
        pairs = DuplicateDetector().detect(frame)
        clusters, members = build_clusters(pairs, frame, "LOW")
        self.assertEqual(len(clusters), 1)
        self.assertEqual(clusters.iloc[0]["SIZE"], 3)
        self.assertEqual(clusters.iloc[0]["SURVIVOR_ID"], "C")  # active beats inactive; then most POs
        self.assertEqual(members.iloc[0]["SUPPLIER_ID"], "C")

    def test_survivor_reason_is_explained(self):
        ranking = survivor_ranking(
            [
                {"SUPPLIER_ID": "1", "SUPPLIER_NAME": "x", "ACTIVE_STATUS": "ACTIVE", "PO_COUNT_12M": "4"},
                {"SUPPLIER_ID": "2", "SUPPLIER_NAME": "x", "ACTIVE_STATUS": "ACTIVE", "PO_COUNT_12M": "1"},
            ]
        )
        self.assertEqual(ranking[0][0]["SUPPLIER_ID"], "1")
        self.assertIn("4 POs", ranking[0][1])

    def test_sample_clusters_group_the_planted_duplicates(self):
        result = scan_suppliers(load_supplier_master(SAMPLE_WORKBOOK))
        self.assertEqual(len(result.clusters), 30)
        self.assertTrue(result.clusters["SIZE"].eq(3).all())
        self.assertTrue(result.queue[result.queue.TYPE.eq("DUPLICATE") & result.queue.PRIORITY.eq("HIGH")].CLUSTER_ID.ne("").all())


class CompareAndReportTests(unittest.TestCase):
    def test_comparison_flags_match_similar_and_masks_bank_accounts(self):
        rows = compare_records(
            {"SUPPLIER_NAME": "Acme Trading Ltd", "TAX_ID": "GB 123", "BANK_ACCOUNT": "GB00BANK12345678", "CITY": "London"},
            {"SUPPLIER_NAME": "ACME Trading", "TAX_ID": "gb123", "BANK_ACCOUNT": "GB00BANK99995678", "CITY": "Leeds"},
        )
        status = {row["FIELD"]: row["STATUS"] for row in rows}
        self.assertEqual(status["SUPPLIER_NAME"], "match")
        self.assertEqual(status["TAX_ID"], "match")
        self.assertEqual(status["BANK_ACCOUNT"], "different")
        self.assertEqual(status["CITY"], "different")
        bank = next(row for row in rows if row["FIELD"] == "BANK_ACCOUNT")
        self.assertTrue(bank["A"].endswith("5678") and "BANK" not in bank["A"])
        self.assertEqual(mask("123"), "123")

    def test_excel_report_has_all_sheets_and_neutralises_formulas(self):
        frame = _frame(
            [
                {"SUPPLIER_ID": "1", "SUPPLIER_NAME": '=HYPERLINK("x")', "TAX_ID": "T1"},
                {"SUPPLIER_ID": "2", "SUPPLIER_NAME": "Acme", "TAX_ID": "T1"},
            ]
        )
        report = load_workbook(BytesIO(build_excel_report(scan_suppliers(frame), [], "test.csv")))
        self.assertEqual(report.sheetnames[:3], ["Summary", "Duplicate clusters", "Cluster members"])
        self.assertIn("Quality findings", report.sheetnames)
        values = [cell.value for row in report["Duplicate pairs"].iter_rows() for cell in row]
        self.assertFalse(any(isinstance(v, str) and v.startswith("=") for v in values))

    def test_cli_scan_writes_queue_report_and_clusters(self):
        with TemporaryDirectory() as temp:
            queue, report, clusters = (str(Path(temp) / name) for name in ("q.csv", "r.xlsx", "c.csv"))
            code = cli_main(["scan", str(EXAMPLE_CSV), "--output", queue, "--report", report, "--clusters", clusters])
            self.assertEqual(code, 0)
            self.assertGreater(len(pd.read_csv(queue)), 0)
            self.assertIn("Summary", load_workbook(report).sheetnames)
            self.assertGreater(len(pd.read_csv(clusters)), 0)


if __name__ == "__main__":
    unittest.main()


class CommunityFeatureTests(unittest.TestCase):
    def test_iban_validation(self):
        from supplier_agent.validators import iban_problem, make_iban

        self.assertIsNone(iban_problem("GB82 WEST 1234 5698 7654 32"))
        self.assertIsNone(iban_problem(make_iban("NL", "ABNA0417164300")))
        self.assertIn("check digits", iban_problem("GB83WEST12345698765432"))
        self.assertIn("length", iban_problem("DE8937040044053201300"))
        self.assertIsNone(iban_problem("123456789"))  # domestic account numbers are not checked
        self.assertIsNone(iban_problem("US12ABC"))  # the US does not use IBANs

    def test_invalid_iban_is_a_high_priority_finding(self):
        from supplier_agent.quality import QualityConfig, check_quality

        frame = _frame([{"SUPPLIER_ID": "1", "SUPPLIER_NAME": "Acme", "BANK_ACCOUNT": "GB83WEST12345698765432"}])
        findings = check_quality(frame)
        row = findings[findings.RULE_CODE.eq("INVALID_IBAN")].iloc[0]
        self.assertEqual(row.PRIORITY, "HIGH")
        self.assertTrue(check_quality(frame, QualityConfig(check_ibans=False)).query("RULE_CODE == 'INVALID_IBAN'").empty)

    def test_international_legal_forms_are_ignored_when_matching(self):
        from supplier_agent.normalize import normalize_company_name as name

        self.assertEqual(name("Northwind Packaging Sp. z o.o."), "northwind packaging")
        self.assertEqual(name("Northwind Packaging SA de CV"), "northwind packaging")
        self.assertEqual(name("Northwind Packaging Pty Ltd"), "northwind packaging")
        self.assertEqual(name("Ltd"), "ltd")  # never reduced to nothing

    def test_sample_planted_issues_are_found(self):
        result = scan_suppliers(load_supplier_master(SAMPLE_WORKBOOK))
        rules = set(result.quality_findings.RULE_CODE)
        for rule in ("INVALID_IBAN", "PLACEHOLDER_IDENTIFIER", "CURRENCY_FORMAT", "COUNTRY_FORMAT", "BANK_SCOPE_MISSING"):
            self.assertIn(rule, rules)

    def test_action_plan_uses_latest_confirmations(self):
        from supplier_agent.actions import action_plan

        frame = _frame(
            [
                {"SUPPLIER_ID": "A", "SUPPLIER_NAME": "Acme Ltd", "TAX_ID": "T1", "ACTIVE_STATUS": "ACTIVE", "PO_COUNT_12M": "40"},
                {"SUPPLIER_ID": "B", "SUPPLIER_NAME": "ACME Limited", "TAX_ID": "T1", "ACTIVE_STATUS": "ACTIVE", "PO_COUNT_12M": "2"},
                {"SUPPLIER_ID": "C", "SUPPLIER_NAME": "Acme", "TAX_ID": "T1", "ACTIVE_STATUS": "INACTIVE"},
            ]
        )
        queue = scan_suppliers(frame).queue
        duplicates = queue[queue.TYPE.eq("DUPLICATE")]
        latest = {
            fid: {"decision": "CONFIRM_DUPLICATE", "reviewer": "QA", "reviewed_at": "2026-01-01", "notes": ""}
            for fid in duplicates.FINDING_ID
        }
        plan = action_plan(queue, latest, frame)
        self.assertEqual(set(plan.KEEP_ID), {"A"})
        self.assertEqual(set(plan.RETIRE_ID), {"B", "C"})
        for row in duplicates.itertuples():
            if "B" in (row.SUPPLIER_ID, row.CANDIDATE_ID):  # a later review overturns B's confirmation
                latest[row.FINDING_ID] = {"decision": "NOT_DUPLICATE"}
        self.assertEqual(set(action_plan(queue, latest, frame).RETIRE_ID), {"C"})


class ContactSignalTests(unittest.TestCase):
    def test_shared_email_or_phone_links_records_across_countries(self):
        frame = _frame(
            [
                {"SUPPLIER_ID": "1", "SUPPLIER_NAME": "Alpha Print Ltd", "EMAIL": "AP@AlphaPrint.example", "PHONE": "+44 20 7946 0001"},
                {"SUPPLIER_ID": "2", "SUPPLIER_NAME": "Alpha Print GmbH", "COUNTRY_CODE": "DE", "EMAIL": "ap@alphaprint.example "},
                {"SUPPLIER_ID": "3", "SUPPLIER_NAME": "Zulu Foods", "PHONE": "020 7946 0001"},
                {"SUPPLIER_ID": "4", "SUPPLIER_NAME": "Kilo Steel", "EMAIL": "noreply@kilo.example"},
                {"SUPPLIER_ID": "5", "SUPPLIER_NAME": "Hotel Paints", "EMAIL": "noreply@kilo.example"},
            ]
        )
        pairs = DuplicateDetector().detect(frame)
        found = {frozenset((row.SUPPLIER_ID_A, row.SUPPLIER_ID_B)): row for row in pairs.itertuples()}
        self.assertTrue(found[frozenset({"1", "2"})].SAME_EMAIL)
        self.assertIn("same email", found[frozenset({"1", "2"})].EVIDENCE)
        self.assertTrue(found[frozenset({"1", "3"})].SAME_PHONE)
        self.assertNotIn(frozenset({"4", "5"}), found)  # no-reply mailboxes are not identity evidence


class TaxIdTests(unittest.TestCase):
    def test_formats_and_check_digits(self):
        from supplier_agent.taxid import tax_id_problem

        valid = [
            ("DE", "DE136695976"),
            ("FR", "FR40303265045"),
            ("BE", "BE0403019261"),
            ("IT", "IT00743110157"),
            ("PL", "PL5260250995"),
            ("AU", "51 824 753 556"),
            ("BR", "11.222.333/0001-81"),
            ("US", "12-3456789"),
            ("GB", "GB 123 4567 89"),
            ("IN", "27AAPFU0939F1ZV"),
            ("CH", "CHE-123.456.789 MWST"),
            ("NL", "NL123456789B01"),
            ("GR", "EL123456789"),
            ("FR", "DE136695976"),
        ]  # a valid German VAT held by a French supplier
        for country, value in valid:
            self.assertIsNone(tax_id_problem(country, value), (country, value))
        self.assertIn("check digit", tax_id_problem("DE", "DE136695977"))
        self.assertIn("check digit", tax_id_problem("BR", "11.222.333/0001-82"))
        self.assertIn("format", tax_id_problem("US", "1234"))
        self.assertIsNone(tax_id_problem("MA", "anything"))  # countries without a rule are not checked
        self.assertIsNone(tax_id_problem("DE", ""))

    def test_tax_id_rule_in_quality_findings(self):
        from supplier_agent.quality import QualityConfig, check_quality

        frame = _frame([{"SUPPLIER_ID": "1", "SUPPLIER_NAME": "Acme", "COUNTRY_CODE": "DE", "TAX_ID": "DE136695977"}])
        self.assertIn("TAX_ID_FORMAT", set(check_quality(frame).RULE_CODE))
        self.assertNotIn("TAX_ID_FORMAT", set(check_quality(frame, QualityConfig(check_tax_ids=False)).RULE_CODE))
