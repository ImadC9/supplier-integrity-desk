"""Formatted Excel report of a scan, ready to share with data owners."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .exports import safe_cell
from .quality import RULE_DESCRIPTIONS

HEADER_FILL = PatternFill("solid", fgColor="0B3558")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PRIORITY_FILLS = {
    "HIGH": PatternFill("solid", fgColor="F8D7D3"),
    "MEDIUM": PatternFill("solid", fgColor="FBEBC8"),
    "LOW": PatternFill("solid", fgColor="E3EEF2"),
}

PAIR_COLUMNS = {
    "SUPPLIER_ID_A": "Supplier A",
    "SUPPLIER_NAME_A": "Name A",
    "SUPPLIER_ID_B": "Supplier B",
    "SUPPLIER_NAME_B": "Name B",
    "COUNTRY_CODE": "Country",
    "LIKELIHOOD": "Priority",
    "DUPLICATE_SCORE": "Rule score",
    "SAME_TAX_ID": "Same tax ID",
    "SAME_ASSOCIATION_NO": "Same registration no.",
    "SAME_BANK_ACCOUNT": "Same bank account",
    "SAME_EMAIL": "Same email",
    "SAME_PHONE": "Same phone",
    "NAME_SIMILARITY": "Name similarity %",
    "ADDRESS_SIMILARITY": "Address similarity %",
    "EVIDENCE": "Evidence",
    "RECOMMENDATION": "Recommendation",
}


def _write(writer: pd.ExcelWriter, frame: pd.DataFrame, sheet: str, priority_column: str | None = None) -> None:
    frame = frame.map(safe_cell)  # neutralise formulas whatever the column dtype
    frame.to_excel(writer, sheet_name=sheet, index=False)
    sheet_obj = writer.sheets[sheet]
    for cell in sheet_obj[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet_obj.freeze_panes = "A2"
    if len(frame.columns):
        sheet_obj.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{max(len(frame) + 1, 1)}"
    for index, column in enumerate(frame.columns, start=1):
        values = [str(column)] + [str(value) for value in frame[column].head(200).tolist()]
        sheet_obj.column_dimensions[get_column_letter(index)].width = min(max(len(value) for value in values) + 2, 60)
    if priority_column and priority_column in frame.columns:
        column_index = list(frame.columns).index(priority_column) + 1
        for row in range(2, len(frame) + 2):
            fill = PRIORITY_FILLS.get(str(sheet_obj.cell(row=row, column=column_index).value))
            if fill:
                sheet_obj.cell(row=row, column=column_index).fill = fill


def build_excel_report(scan, reviews: list[dict] | None = None, source_name: str = "", actions: pd.DataFrame | None = None) -> bytes:
    """Return an .xlsx with summary, clusters, pairs, quality findings and reviews."""
    reviews = reviews or []
    pairs = scan.duplicate_pairs
    quality = scan.quality_findings
    summary_rows = [
        ("Generated (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M")),
        ("Source", source_name or "—"),
        ("Dataset fingerprint", scan.dataset_id[:16]),
        ("Suppliers scanned", scan.supplier_count),
        ("Duplicate pairs — high", int(pairs["LIKELIHOOD"].eq("HIGH").sum()) if not pairs.empty else 0),
        ("Duplicate pairs — medium", int(pairs["LIKELIHOOD"].eq("MEDIUM").sum()) if not pairs.empty else 0),
        ("Duplicate pairs — low", int(pairs["LIKELIHOOD"].eq("LOW").sum()) if not pairs.empty else 0),
        ("Duplicate clusters", len(scan.clusters)),
        ("Suppliers inside clusters", int(scan.clusters["SIZE"].sum()) if not scan.clusters.empty else 0),
        ("Data-quality findings", len(quality)),
        ("Decisions recorded", len(reviews)),
        ("Records to block (action plan)", 0 if actions is None else len(actions)),
        ("", ""),
        (
            "Note",
            "Rule scores prioritise review; they are not probabilities. Suggested survivors are advisory. No supplier record was changed.",
        ),
    ]
    rule_rows = []
    if not quality.empty:
        counts = quality.groupby(["RULE_CODE", "PRIORITY"]).size().reset_index(name="FINDINGS")
        for item in counts.itertuples(index=False):
            rule_rows.append(
                {
                    "Rule": item.RULE_CODE,
                    "Priority": item.PRIORITY,
                    "Findings": item.FINDINGS,
                    "Meaning": RULE_DESCRIPTIONS.get(item.RULE_CODE, ""),
                }
            )

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _write(writer, pd.DataFrame(summary_rows, columns=["Item", "Value"]), "Summary")
        _write(writer, scan.clusters.rename(columns=str.title), "Duplicate clusters", "Top_Likelihood")
        _write(writer, scan.cluster_members.rename(columns=str.title), "Cluster members")
        _write(
            writer,
            pairs[list(PAIR_COLUMNS)].rename(columns=PAIR_COLUMNS)
            if not pairs.empty
            else pd.DataFrame(columns=list(PAIR_COLUMNS.values())),
            "Duplicate pairs",
            "Priority",
        )
        _write(writer, quality.rename(columns=str.title), "Quality findings", "Priority")
        _write(writer, pd.DataFrame(rule_rows, columns=["Rule", "Priority", "Findings", "Meaning"]), "Quality rules", "Priority")
        _write(writer, scan.completeness.rename(columns=str.title), "Field completeness")
        if actions is not None and not actions.empty:
            _write(writer, actions.rename(columns=str.title), "Action plan")
        review_frame = pd.DataFrame(reviews).drop(columns=["case_json"], errors="ignore")
        _write(writer, review_frame, "Decisions")
    return buffer.getvalue()
